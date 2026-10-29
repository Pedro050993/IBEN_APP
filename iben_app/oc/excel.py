"""Excel da extracao de OCs: Resumo, itens da obra, Cruzamento memoria, Pastas e status."""
import datetime as dt
import os

import openpyxl

from iben_app.comum import COR_AMARELO, COR_LARANJA, COR_STATUS, estilo, tabela

COLS_ITENS = ["Status", "Tarefa", "Nº pedido", "Nº OC", "Data OC", "Fornecedor", "Item", "Descrição",
              "Unid.", "Qtde", "R$ unit.", "R$ total", "Conferência", "Pasta", "CNPJ", "Observação"]
COLS_CRUZ = ["Nº OC", "Data OC", "Pasta", "Fornecedor", "Total da OC (R$)", "Método", "Confiança",
             "Linha(s) na memória", "EAP", "Item na memória", "Fornecedor na memória",
             "Valor na memória (R$)", "Diferença (R$)", "Citada em outras linhas", "Situação"]


def _data(s):
    try:
        return dt.datetime.strptime(s, "%d/%m/%Y")
    except (TypeError, ValueError):
        return None


def gerar(obra, procs, cruz, saida, memoria=None, pasta_origem=None):
    e = estilo()
    por_pasta = {id(c["pr"]): c for c in cruz}
    wb = openpyxl.Workbook()
    res = wb.active
    res.title = "Resumo"
    aba_itens = f"Itens {obra}"[:31]
    ws = wb.create_sheet(aba_itens)
    linhas, marcas = [], []
    for pr in procs:
        c = por_pasta.get(id(pr))
        cab = pr.get("cab") or {}
        tarefa_mem, cor_tarefa, obs_extra = "", None, ""
        if c and c["metodo"] == "Nº pedido" and c["eaps"]:
            tarefa_mem = "; ".join(c["eaps"])
        elif c and c["metodo"] == "Semelhança" and c["confianca"] in ("Alta", "Média") and c["eaps"]:
            tarefa_mem, cor_tarefa = c["eaps"][0], COR_LARANJA
            ln = c["linhas"][0]
            obs_extra = f"EAP por semelhança com a linha {ln['linha']} da memória, validar"
        obs_pasta = " | ".join(pr.get("obs", []))
        base = [pr.get("status", ""), None, cab.get("n_pedido"), pr.get("n_oc"), _data(cab.get("data")),
                cab.get("fornecedor")]
        if not pr.get("itens"):
            linhas.append(base[:1] + [tarefa_mem or None] + base[2:] + [None] * 7 + [pr["pasta"], cab.get("cnpj"),
                                                                                 " | ".join(x for x in (obs_pasta, obs_extra) if x)])
            marcas.append(cor_tarefa if tarefa_mem else COR_AMARELO)
            continue
        for it in pr["itens"]:
            tarefa = it.get("tarefa") or tarefa_mem or None
            obs = " | ".join(x for x in (it.get("obs"), obs_pasta, obs_extra if not it.get("tarefa") else "") if x)
            linhas.append([pr.get("status", ""), tarefa] + base[2:] +
                          [it["item"], it["descricao"], it["unid"], it["qtde"], it["vu"], it["vt"], None,
                           pr["pasta"], cab.get("cnpj"), obs])
            marcas.append(None if it.get("tarefa") else (cor_tarefa if tarefa else COR_AMARELO))
    ws.append([f"OCs da obra 400.{obra} | Processo de Compras"])
    ws["A1"].font = e["titulo"]
    ws.append([f"Origem: {pasta_origem or ''} | gerado em {dt.datetime.now():%d/%m/%Y %H:%M} | valores em R$"])
    ws["A2"].font = e["sub"]
    r0, r1 = tabela(ws, COLS_ITENS, linhas, [11, 12, 12, 18, 11, 30, 6, 50, 7, 11, 13, 14, 12, 40, 19, 50], linha0=4,
                    formatos={4: "dd/mm/yyyy", 9: "#,##0.0000", 10: "#,##0.00", 11: "#,##0.00", 12: "#,##0.00"})
    for i, cor in enumerate(marcas):
        r = r0 + i
        ws.cell(r, 13).value = f'=IF(OR(J{r}="",K{r}=""),"",ROUND(J{r}*K{r}-L{r},2))'
        st = ws.cell(r, 1).value
        if st in COR_STATUS:
            ws.cell(r, 1).fill = e["fill"](COR_STATUS[st])
        if cor:
            ws.cell(r, 2).fill = e["fill"](cor)
        ws.cell(r, 8).alignment = ws.cell(r, 16).alignment = e["quebra"]

    # Cruzamento memoria
    wc = wb.create_sheet("Cruzamento memória")
    lc = []
    for c in cruz:
        pr, ln0 = c["pr"], (c["linhas"][0] if c["linhas"] else {})
        lc.append([pr.get("n_oc"), _data((pr.get("cab") or {}).get("data")), pr["pasta"],
                   (pr.get("cab") or {}).get("fornecedor"), (pr.get("totais") or {}).get("total"), c["metodo"],
                   c["confianca"], ", ".join(str(x["linha"]) for x in c["linhas"]), "; ".join(c["eaps"]),
                   ln0.get("item"), ln0.get("fornecedor"), c["valor_memoria"], c["diferenca"],
                   ", ".join(map(str, c["citada"])), c["situacao"]])
    lc.sort(key=lambda x: -abs(x[12] or 0))  # maiores diferencas primeiro
    tabela(wc, COLS_CRUZ, lc, [18, 11, 40, 28, 15, 18, 10, 14, 14, 40, 26, 15, 14, 14, 45],
           formatos={1: "dd/mm/yyyy", 4: "#,##0.00", 11: "#,##0.00", 12: "#,##0.00"})
    for r in range(2, len(lc) + 2):
        if wc.cell(r, 6).value == "Semelhança":
            wc.cell(r, 9).fill = e["fill"](COR_LARANJA)

    # Pastas e status
    wp = wb.create_sheet("Pastas e status")
    lp = [[pr["pasta"], pr.get("n_oc"), pr.get("status"), pr["modificada"],
           os.path.basename(pr["pdf"]) if pr.get("pdf") else None, len(pr.get("itens", [])),
           (pr.get("totais") or {}).get("total"), " | ".join(pr.get("obs", []))] for pr in procs]
    tabela(wp, ["Pasta", "Nº OC", "Status", "Modificada em", "PDF lido", "Itens", "Total (R$)", "Observação"], lp,
           [45, 18, 13, 16, 40, 7, 15, 70], formatos={3: "dd/mm/yyyy hh:mm", 6: "#,##0.00"})
    for r in range(2, len(lp) + 2):
        st = wp.cell(r, 3).value
        if st in COR_STATUS:
            wp.cell(r, 3).fill = e["fill"](COR_STATUS[st])

    # Resumo com formulas apontando para as abas
    q = f"'{aba_itens}'"
    res.append([f"Extração de OCs | obra 400.{obra}"])
    res["A1"].font = e["titulo"]
    res.append([f"Gerado em {dt.datetime.now():%d/%m/%Y %H:%M}. Janela: pastas modificadas nos últimos 3 meses. "
                "Candidata por semelhança é hipótese, valide com o engenheiro."])
    res["A2"].font = e["sub"]
    if memoria:
        res.append([f"Memória: {os.path.basename(memoria['arquivo'])} | aba {memoria['aba']}"])
        res["A3"].font = e["sub"]
    itens = [
        ("Pastas na janela", "=COUNTA('Pastas e status'!A:A)-1", None),
        ("Pastas OK", "=COUNTIF('Pastas e status'!C:C,\"OK\")", None),
        ("Pastas para verificar", "=COUNTIF('Pastas e status'!C:C,\"Verificar\")", None),
        ("Canceladas", "=COUNTIF('Pastas e status'!C:C,\"Cancelado\")", None),
        ("Sem OC", "=COUNTIF('Pastas e status'!C:C,\"Sem OC\")", None),
        ("Fora do escopo", "=COUNTIF('Pastas e status'!C:C,\"Fora do escopo\")", None),
        ("Linhas de itens", f"=COUNTA({q}!H5:H5000)", None),
        ("Valor total dos itens (R$)", f"=SUM({q}!L:L)", "#,##0.00"),
        ("Itens com Qtde x Unit diferente do total", f"=SUMPRODUCT(--(ABS(N({q}!M5:M5000))>0.01))", None),
        ("Itens sem Tarefa", f"=COUNTIFS({q}!B5:B5000,\"\",{q}!H5:H5000,\"<>\")", None),
        ("Cruzamento: casadas por Nº pedido", "=COUNTIF('Cruzamento memória'!F:F,\"Nº pedido\")", None),
        ("Cruzamento: candidatas por semelhança", "=COUNTIF('Cruzamento memória'!F:F,\"Semelhança\")", None),
        ("Cruzamento: só citadas", "=COUNTIF('Cruzamento memória'!F:F,\"Só citada\")", None),
        ("Cruzamento: sem correspondência", "=COUNTIF('Cruzamento memória'!F:F,\"Sem correspondência\")", None),
        ("Erros de fórmula no arquivo", f"=SUMPRODUCT(--ISERROR({q}!M5:M5000))", None),
    ]
    r = 5
    res.cell(r, 1, "Indicador").font = e["cab"]
    res.cell(r, 2, "Valor").font = e["cab"]
    res.cell(r, 1).fill = res.cell(r, 2).fill = e["fundo_cab"]
    for nome, f, fmt in itens:
        r += 1
        res.cell(r, 1, nome).font = e["tinta"]
        c = res.cell(r, 2, f)
        c.font = e["negrito"]
        if fmt:
            c.number_format = fmt
    res.column_dimensions["A"].width = 44
    res.column_dimensions["B"].width = 18
    if not memoria:
        wc.cell(2, 1, "Memória de cálculo não informada: cruzamento não rodou.")
    wb.save(saida)
    return saida
