"""Arquivos sinteticos que imitam o layout das planilhas e PDFs da Iben (para teste).
Os casos de regressao reais (Vila Flora II: OC 00351 -> linha 1266, 00576 -> 986, 00349 por
Nº pedido) precisam dos arquivos reais e entram no teste quando o caminho for informado."""
import datetime as dt
import os

import openpyxl
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas


def pdf_texto(path, linhas):
    c = canvas.Canvas(path, pagesize=A4)
    c.setFont("Courier", 7.5)
    y = 810
    for ln in linhas:
        c.drawString(20, y, ln)
        y -= 10
    c.save()


def oc_pdf(path, n_oc, fornecedor, itens, frete=0.0, desconto=0.0, n_pedido="000001565",
           data="15/09/2026", subtotal=None, tarefa=False):
    def br(v):
        return f"{v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    L = ["IBEN ENGENHARIA                                   PMO.01.01 - ORDEM DE COMPRA",
         f"OC: {n_oc}                                  Nº PEDIDO: {n_pedido}",
         f"DATA: {data}",
         f"FORNECEDOR: {fornecedor}              CNPJ: 12.345.678/0001-90",
         "COMPRADOR: FULANO",
         "",
         ("ITEM  TAREFA   " if tarefa else "ITEM  ") + "DESCRIÇÃO                              UNID.   QTDE      R$ UNIT.     R$ TOTAL"]
    soma = 0.0
    for i, it in enumerate(itens, 1):
        desc, un, q, vu = it[:4]
        vt = it[4] if len(it) > 4 and it[4] is not None else round(q * vu, 2)
        soma += vt
        extra = it[5] if len(it) > 5 else []
        tar = f"{it[6]:8s} " if tarefa and len(it) > 6 else ""
        L.append(f"{i:<5d} {tar}{desc:38s} {un:6s} {br(q):>8s} {br(vu):>12s} {br(vt):>12s}")
        L.extend("      " + x for x in extra)
        L.append(f"      OBS: entrega na obra item {i}")
    L.append("")
    sub = soma if subtotal is None else subtotal
    L.append(f"SUBTOTAL: R$ {br(sub)}")
    if frete:
        L.append(f"FRETE: + R$ {br(frete)}")
    if desconto:
        L.append(f"DESCONTO: - R$ {br(desconto)}")
    L.append(f"TOTAL: R$ {br(sub + frete - desconto)}")
    L.append("CONDIÇÃO DE PAGAMENTO: 28 DDL. Observação: o valor total inclui impostos.")
    pdf_texto(path, L)
    return round(sub + frete - desconto, 2)


def pasta_compras(raiz, hoje=None):
    """Pasta Processo de Compras da obra 400.137 com casos tipicos."""
    hoje = hoje or dt.datetime.now()
    os.makedirs(raiz, exist_ok=True)
    casos = {}

    def proc(nome, dias=10):
        p = os.path.join(raiz, nome)
        os.makedirs(p, exist_ok=True)
        t = (hoje - dt.timedelta(days=dias)).timestamp()
        return p, t

    p, t = proc("400.137.2.00349 - Concreto usinado (Mov 1201)")
    casos["00349"] = oc_pdf(os.path.join(p, "OC 400.137.2.00349.pdf"), "400.137.2.00349", "CONCRETEX LTDA",
                            [("Concreto fck 30 MPa bombeado", "m3", 120, 520.0)], frete=350.0)
    oc_pdf(os.path.join(p, "OC 400.137.2.00349_R01.pdf"), "400.137.2.00349", "CONCRETEX LTDA",
           [("Concreto fck 30 MPa bombeado", "m3", 120, 520.0)], frete=350.0)
    casos["00349"] = 120 * 520 + 350
    os.utime(p, (t, t))

    p, t = proc("400.137.2.00351 - MO Pintura Fase 3 (Mov 1300)")
    casos["00351"] = oc_pdf(os.path.join(p, "OC 400.137.2.00351.pdf"), "400.137.2.00351", "PINTURAS SILVA ME",
                            [("Mão de obra pintura interna", "m2", 3500, 18.5, None,
                              ["paredes e tetos, fase 3"]), ("Mão de obra pintura fachada", "m2", 1200.5, 22.0)])
    os.utime(p, (t, t))

    p, t = proc("400.137.2.00576 - Esquadrias de alumínio (Mov 1410)")
    casos["00576"] = oc_pdf(os.path.join(p, "PC 400.137.2.00576.pdf"), "400.137.2.00576", "ALUMIX ESQUADRIAS LTDA",
                            [("Janela de correr alumínio 1,20x1,00", "un", 96, 1450.0, 139200.03)])
    os.utime(p, (t, t))

    p, t = proc("400.137.2.00600 - Paisagismo (Mov 1500)")
    casos["00600"] = oc_pdf(os.path.join(p, "OC 400.137.2.00600.pdf"), "400.137.2.00600", "VERDE VIVO",
                            [("Plantio de grama esmeralda", "m2", 800, 25.0)], subtotal=21000.0)
    os.utime(p, (t, t))

    p, t = proc("400.137.2.00610 - Locação de andaime CANCELADO")
    os.utime(p, (t, t))
    p, t = proc("400.137.2.00611 - Caçamba")
    os.utime(p, (t, t))
    p, t = proc("400.139.2.00010 - Pedido de outra obra")
    os.utime(p, (t, t))
    p, t = proc("400.137.2.00100 - Antigo", dias=200)
    os.utime(p, (t, t))
    p, t = proc("Processos Antigos")
    os.utime(p, (t, t))
    return casos


def memoria_xlsx(path, extra_abas=True):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Memória de Cálculo - Previsão"
    ws["B2"], ws["C2"] = "CONTINGÊNCIA", None
    ws["B3"], ws["C3"] = "ATUAL", 1000000.0
    cab = ["", "EAP", "MAT/MO/EQUIP", "ITEM", "FORNECEDOR", "OBSERVAÇÕES", "N° PEDIDO", "VALOR TOTAL DO CONTRATO",
           "INCORRIDO", "À PAGAR", "RETENÇÃO", "SALDO"]
    for j, h in enumerate(cab, 1):
        ws.cell(7, j, h)
    linhas = [
        ("4.1", "MAT", "Concreto usinado fck 30", "Concretex", "", "400137200349", 62750.0),
        ("5.9.2", "MO", "MO pintura interna e fachada fase 3", "Pinturas Silva", "", "Previsão", 90000.0),
        ("5.9.2", "MO", "MO pintura fase 2", "Pinturas Silva", "", "400.137.2.00300", 50000.0),
        ("5.8.4", "MAT", "Deca anel", "Deca", "", "4001370351", 1200.0),
        ("5.4.1", "MAT", "Esquadrias de alumínio janelas", "Alumix", "ver OC 0576 e 0577", "Previsão", 140000.0),
        ("4.7", "MO", "Paisagismo", "", "", "Previsão", 15000.0),
        ("1.2", "EQUIP", "Caçamba entulho", "", "", "Previsão", 5000.0),
    ]
    for i, (eap, t, item, forn, obs, ped, val) in enumerate(linhas, 8):
        for j, v in enumerate((eap, t, item, forn, obs, ped, val), 2):
            ws.cell(i, j, v)
        ws.cell(i, 9, val * 0.4)
        ws.cell(i, 10, val * 0.1)
        ws.cell(i, 11, f"=I{i}*0.05")
        ws.cell(i, 12, f"=H{i}-I{i}-J{i}")
    ws.cell(10, 12, 999.0)  # valor digitado onde deveria haver formula (auditoria)
    ws.cell(10, 11, f"=I10*0.05+L10")
    wb.save(path)
    return path


def analise_custos(pasta, nome="Análise de Custos - VILLA FLORA II - Out_26_R01.xlsx", mes=dt.date(2026, 10, 1),
                   recalc=True):
    """Analise de Custos sintetica: Analise Financeira (bloco ANALISE PREVISAO - 100%), Desvios,
    Memoria, Incorrido, Retencoes e A pagar. Recalculada no LibreOffice para ter valores em cache."""
    from openpyxl.styles import PatternFill
    from tests.util import recalcular
    os.makedirs(pasta, exist_ok=True)
    base = os.path.join(pasta, "_base_" + nome)
    memoria_xlsx(base)
    wb = openpyxl.load_workbook(base)
    af = wb.create_sheet("Análise Financeira", 0)
    af["G5"], af["G6"] = "MÊS ANÁLISE", dt.datetime(mes.year, mes.month, 1)
    af["J8"] = "ANÁLISE PREVISÃO - 100%"
    cab = {2: "EAP", 3: "DESCRIÇÃO", 5: "ORÇAMENTO", 10: "DIFERENÇA", 11: "DIFERENÇA MÊS ANTERIOR",
           12: "VARIAÇÃO (INCC - MÊS)", 13: "VARIAÇÃO", 14: "JUSTIFICATIVA DO DESVIO"}
    for c, h in cab.items():
        af.cell(9, c, h)
    # eap, descricao, orcamento, dif atual, dif anterior, incc, justificativa
    folhas = {
        "1.1": ("Canteiro", 500000, -20000, -15000, -1200, "AUMENTO DO PRAZO DO CANTEIRO E LOCAÇÃO DE CONTAINER"),
        "1.2": ("Equipamentos", 300000, 10000, 10300, -300, "VARIAÇÃO DE INCC"),
        "2.1": ("Estrutura", 2000000, 50000, 52000, -2000, "GANHO INCC"),
        "2.2": ("Alvenaria", 800000, -8000, -10500, -500, "NOVO PEDIDO OC 400.137.2.00576 ABAIXO DA PREVISÃO"),
        "2.3": ("Pintura", 400000, 0, 0, 0, ""),
    }
    r = 10
    pais = {}
    for pai, desc in (("1", "SERVIÇOS PRELIMINARES"), ("2", "CONSTRUÇÃO CIVIL")):
        r0 = r
        af.cell(r, 2, pai)
        af.cell(r, 3, desc)
        r += 1
        for eap, (d, orc, dif, ant, incc, just) in folhas.items():
            if eap.split(".")[0] != pai:
                continue
            for c, v in zip((2, 3, 5, 10, 11, 12, 14), (eap, d, orc, dif, ant, incc, just)):
                af.cell(r, c, v)
            af.cell(r, 13, f"=J{r}-K{r}")
            r += 1
        for c in (5, 10, 11, 12, 13):
            L = openpyxl.utils.get_column_letter(c)
            af.cell(r0, c, f"=SUBTOTAL(9,{L}{r0 + 1}:{L}{r - 1})")
        pais[pai] = r0
    af.cell(r, 3, "CONSTRUÇÃO")
    for c in (5, 10, 11, 12, 13):
        L = openpyxl.utils.get_column_letter(c)
        af.cell(r, c, f"=SUBTOTAL(9,{L}10:{L}{r - 1})")
    tot = r
    r += 2
    for c, v in zip((2, 3, 5, 10, 11, 12, 14), ("5.13.3", "Muro de arrimo", 90000, -5000, -4000, -100, "SM MURO EXTERNO")):
        af.cell(r, c, v)
    af.cell(r, 13, f"=J{r}-K{r}")
    # Desvios
    dv = wb.create_sheet("Desvios", 1)
    dv["B3"] = "RESULTADO DA ANÁLISE DO PREVISÃO - 100%"
    for c, h in {2: "EAP", 3: "DESCRIÇÃO", 4: "DIFERENÇA", 5: "VARIAÇÃO - INCC", 6: "VARIAÇÃO"}.items():
        dv.cell(4, c, h).fill = PatternFill("solid", fgColor="1F3864")
    linhas = [("2.1", "Estrutura", 50000, -2000, -2000), ("1.2", "Equipamentos", 10000, -300, -300),
              ("2.3", "Pintura", 0, 0, 0), ("1.1", "Canteiro", -20000, -1200, -5000),
              ("2.2", "Alvenaria", -8000, -500, 2500)]
    for i, ln in enumerate(linhas, 5):
        for c, v in enumerate(ln, 2):
            dv.cell(i, c, v)
    rf = 5 + len(linhas)
    dv.cell(rf, 2, "RESULTADO FINAL")
    dv.cell(rf, 4, f"='Análise Financeira'!J{tot}")
    dv.cell(rf, 6, f"='Análise Financeira'!M{tot}")
    dv.cell(rf, 2).fill = PatternFill("solid", fgColor="FFF2CC")
    # Incorrido (N-1), Retencoes, A pagar
    inc = wb.create_sheet("Incorrido")
    inc.append(["DATA", "FORNECEDOR", "MÊS", "VALOR BAIXADO", "CUSTO QUE NÃO É DA OBRA"])
    inc.append([dt.datetime(2026, 9, 5), "A", dt.datetime(2026, 9, 1), 100000.0, 0])
    inc.append([dt.datetime(2026, 9, 9), "B", dt.datetime(2026, 9, 1), 50000.0, 2000.0])
    inc.append([dt.datetime(2026, 8, 9), "C", dt.datetime(2026, 8, 1), 70000.0, 0])
    ret = wb.create_sheet("Retenções")
    ret.append(["FORNECEDOR", "RETENÇÃO ATUAL"])
    ret.append(["A", 1234.5])
    ap = wb.create_sheet("A pagar")
    ap.append(["FORNECEDOR", "VALOR BAIXADO"])
    ap.append(["A", 4321.0])
    wb.save(base)
    if not recalc:
        return base
    rec = recalcular(base, pasta)
    final = os.path.join(pasta, nome)
    os.replace(rec, final)
    os.remove(base)
    return final
