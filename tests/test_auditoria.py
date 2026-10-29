import openpyxl

from iben_app.custo import auditoria
from tests.fixtures import memoria_xlsx


def test_auditoria_memoria(tmp_path):
    p = memoria_xlsx(str(tmp_path / "mem.xlsx"))
    wb = openpyxl.load_workbook(p)
    ws = wb["Memória de Cálculo - Previsão"]
    ws["I15"] = "=SUM(I9:I14)"  # bloco comeca na 8: subtotal desalinhado
    ws["H16"], ws["I16"] = 5000.0, 2000.0  # linha sem EAP no fim
    ws["L17"] = "=FALSE()"
    wb.save(p)
    ach = auditoria.auditar(p)
    tipos = {(a["tipo"], a["celula"]) for a in ach}
    assert ("Valor digitado", "L10") in tipos
    assert ("+L parcial", "K10") in tipos
    assert ("Subtotal desalinhado", "I15") in tipos
    assert ("Sem EAP", "B16") in tipos
    sub = next(a for a in ach if a["tipo"] == "Subtotal desalinhado")
    assert sub["correcao"] == "=SUM(I8:I14)"
    saida = auditoria.gerar_excel(p, ach, str(tmp_path / "aud.xlsx"))
    assert openpyxl.load_workbook(saida)["Auditoria"]["A4"].value == "Aba"
