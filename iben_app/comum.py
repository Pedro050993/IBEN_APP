"""Funcoes comuns: caminho informado pelo engenheiro, varredura de pasta, escolha do
arquivo mais recente pelo nome e estilo dos Excel de saida (padrao visual Iben)."""
import datetime as dt
import os
import re
import unicodedata
from urllib.parse import unquote

from iben_app.igile.igile import escolher, mes_do_nome, revisao_do_nome  # noqa: F401  (regra unica do Pedro)


def norm(s):
    """Minusculas, sem acento, espacos (inclusive U+00A0) colapsados."""
    if s is None:
        return ""
    s = unicodedata.normalize("NFKD", str(s).replace("\u00a0", " ")).encode("ascii", "ignore").decode()
    return " ".join(s.lower().split())


def num(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def br(v, casas=2):
    """1234567.8 -> '1.234.567,80'."""
    if v is None:
        return ""
    s = f"{v:,.{casas}f}"
    return s.replace(",", "X").replace(".", ",").replace("X", ".")


def valor_br(s):
    """'1.234,56' -> 1234.56; None se nao for numero."""
    if s is None:
        return None
    s = str(s).strip().replace("R$", "").replace(" ", "")
    if not re.fullmatch(r"-?[\d\.]*,\d+|-?\d+", s):
        return None
    return float(s.replace(".", "").replace(",", "."))


class CaminhoInvalido(ValueError):
    pass


def resolver_caminho(texto):
    """Aceita o caminho como o engenheiro copia: com aspas, file:///, barra invertida ou
    link do SharePoint. Link web nao e lido direto: o app trabalha na pasta sincronizada."""
    if not texto or not texto.strip():
        raise CaminhoInvalido("informe o caminho da pasta ou do arquivo")
    p = texto.strip().strip('"').strip("'")
    if re.match(r"https?://", p, re.I):
        raise CaminhoInvalido(
            "isso é um link da web. Sincronize a pasta no OneDrive (botão Sincronizar ou "
            "Adicionar atalho em Meus arquivos) e cole o caminho do Explorador de Arquivos, "
            "ex.: C:\\Users\\<voce>\\Iben Engenharia\\...\\1. PROCESSO DE COMPRAS")
    if p.lower().startswith("file:"):
        p = unquote(re.sub(r"^file:/*", "", p, flags=re.I))
        if not re.match(r"[A-Za-z]:", p):
            p = "/" + p
    p = os.path.expanduser(p)
    if not os.path.exists(p):
        raise CaminhoInvalido(f"caminho não encontrado: {p}")
    return os.path.normpath(p)


def listar_arquivos(pasta, extensoes=None, recursivo=True):
    """Arquivos da pasta (sem temporarios ~$), com data de modificacao."""
    out = []
    for raiz, dirs, arqs in os.walk(pasta):
        for a in arqs:
            if a.startswith("~$"):
                continue
            if extensoes and not a.lower().endswith(tuple(extensoes)):
                continue
            p = os.path.join(raiz, a)
            out.append({"arquivo": p, "nome": a,
                        "modificado": dt.datetime.fromtimestamp(os.path.getmtime(p))})
        if not recursivo:
            break
    return out


def ultima_planilha(pasta_ou_arquivo, tipo="custo", extensoes=(".xlsm", ".xlsx"), ignorar=()):
    """Escolhe pelo nome (ultimo mes, maior revisao) com a regra do igile.py.
    Retorna o dict do `escolher` mais a lista de arquivos ignorados."""
    if os.path.isfile(pasta_ou_arquivo):
        return {"escolha": {"arquivo": pasta_ou_arquivo, "mes": mes_do_nome(pasta_ou_arquivo)[0],
                            "revisao": revisao_do_nome(pasta_ou_arquivo), "alerta": None,
                            "copia": norm(os.path.basename(pasta_ou_arquivo)).startswith("copia de")},
                "candidatos": []}
    arqs = [a["arquivo"] for a in listar_arquivos(pasta_ou_arquivo, extensoes)
            if not any(norm(x) in norm(a["nome"]) for x in ignorar)]
    return escolher(arqs, tipo)


# ---------------- Excel de saida ----------------

COR_CAB = "3E5C50"       # verde escuro sobrio do log do iGile
COR_TINTA = "2B2B2B"
COR_ALERTA = "F6E7C1"    # ambar claro
COR_LARANJA = "FBE3CF"   # candidata por semelhanca
COR_AMARELO = "FFF2B3"   # celula vazia que pede atencao
COR_STATUS = {"OK": "DCEBDD", "Verificar": "F6E7C1", "Cancelado": "D9D9D9",
              "Sem OC": "F4CCCC", "Fora do escopo": "EDEDED"}


def estilo():
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    return {
        "titulo": Font(name="Calibri", size=14, bold=True, color=COR_CAB),
        "sub": Font(name="Calibri", size=10, italic=True, color="6B6B6B"),
        "cab": Font(name="Calibri", size=10, bold=True, color="FFFFFF"),
        "fundo_cab": PatternFill("solid", fgColor=COR_CAB),
        "tinta": Font(name="Calibri", size=10, color=COR_TINTA),
        "negrito": Font(name="Calibri", size=10, bold=True, color=COR_TINTA),
        "linha": Border(bottom=Side(style="hair", color="B8B8B8")),
        "quebra": Alignment(wrap_text=True, vertical="top"),
        "fill": lambda cor: PatternFill("solid", fgColor=cor),
    }


def tabela(ws, cabecalho, linhas, larguras=None, linha0=1, formatos=None):
    """Escreve cabecalho + linhas com filtro, painel congelado e formatos por coluna
    (formatos = {indice_coluna_0: '#,##0.00'})."""
    from openpyxl.utils import get_column_letter as L
    e = estilo()
    for j, h in enumerate(cabecalho, 1):
        c = ws.cell(linha0, j, h)
        c.font, c.fill = e["cab"], e["fundo_cab"]
        c.alignment = e["quebra"]
    for i, lin in enumerate(linhas, linha0 + 1):
        for j, v in enumerate(lin, 1):
            c = ws.cell(i, j, v)
            c.font, c.border = e["tinta"], e["linha"]
            if formatos and (j - 1) in formatos:
                c.number_format = formatos[j - 1]
    ult = linha0 + max(len(linhas), 1)
    ws.auto_filter.ref = f"A{linha0}:{L(len(cabecalho))}{ult}"
    ws.freeze_panes = ws.cell(linha0 + 1, 1)
    for j, w in enumerate(larguras or [], 1):
        ws.column_dimensions[L(j)].width = w
    return linha0 + 1, ult
