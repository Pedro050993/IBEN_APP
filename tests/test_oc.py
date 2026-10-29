import datetime as dt
import os

import openpyxl
import pytest

from iben_app.oc import cruzamento, excel, extracao
from tests.fixtures import memoria_xlsx, pasta_compras
from tests.util import recalcular


@pytest.fixture(scope="module")
def caso(tmp_path_factory):
    d = tmp_path_factory.mktemp("oc")
    pasta_compras(str(d / "compras"))
    memoria_xlsx(str(d / "mem.xlsx"))
    obra, procs = extracao.extrair_pasta(str(d / "compras"))
    mem = cruzamento.ler_memoria(str(d / "mem.xlsx"))
    cz = cruzamento.cruzar(procs, mem, obra)
    return d, obra, procs, mem, cz


def _pr(procs, seq):
    return next(p for p in procs if p.get("seq") == seq)


def _cz(cz, seq):
    return next(c for c in cz if c["pr"].get("seq") == seq)


def test_obra_e_janela(caso):
    _, obra, procs, _, _ = caso
    assert obra == "137"
    nomes = [p["pasta"] for p in procs]
    assert not any("Antigo" in n for n in nomes), "pasta fora da janela de 3 meses entrou"
    assert _pr(procs, 610)["status"] == "Cancelado"
    assert _pr(procs, 611)["status"] == "Sem OC"
    assert next(p for p in procs if p["pasta"].startswith("400.139"))["status"] == "Fora do escopo"


def test_ultima_revisao_e_totais(caso):
    _, _, procs, _, _ = caso
    p = _pr(procs, 349)
    assert p["pdf"].endswith("_R01.pdf")
    assert p["totais"]["total"] == 62750.0 and p["totais"]["frete"] == 350.0
    assert p["status"] == "OK"


def test_descricao_quebrada_e_arredondamento(caso):
    _, _, procs, _, _ = caso
    it = _pr(procs, 351)["itens"]
    assert it[0]["descricao"].endswith("paredes e tetos, fase 3")
    assert it[1]["qtde"] == 1200.5
    esq = _pr(procs, 576)
    assert esq["status"] == "OK"
    assert "arredondado" in esq["itens"][0]["obs"]


def test_subtotal_divergente_vira_verificar(caso):
    _, _, procs, _, _ = caso
    p = _pr(procs, 600)
    assert p["status"] == "Verificar"
    assert any("SUBTOTAL" in x for x in p["pendencias"])


def test_numeros_na_celula():
    f = cruzamento.numeros_na_celula
    assert f("400.137.2.00351", "137") == [(351, False)]
    assert f("400137200351", "137") == [(351, False)]
    assert f("40013720491", "137") == [(491, False)]
    assert f("0400.137.2.00527", "137") == [(527, False)]
    assert f("400.137.0484", "137") == [(484, False)]
    assert f("4001370351", "137") == []  # 10 digitos: outro pedido, nunca juntar
    assert f("4001372001330", "137") == [(133, True)]
    assert f("400.137.2.00351 / 400.137.2.00352.1", "137") == [(351, False), (352, False)]


def test_cruzamento_camadas(caso):
    _, _, _, _, cz = caso
    c = _cz(cz, 349)
    assert c["metodo"] == "Nº pedido" and c["situacao"] == "valor bate" and c["eaps"] == ["4.1"]
    c = _cz(cz, 351)
    assert c["metodo"] == "Semelhança" and c["confianca"] in ("Alta", "Média")
    assert c["linhas"][0]["linha"] == 9, "deve casar com a fase 3, nao com a 4001370351 (Deca)"
    c = _cz(cz, 576)
    assert c["metodo"] == "Semelhança" and c["linhas"][0]["eap"] == "5.4.1" and 12 in c["citada"]
    c = _cz(cz, 600)
    assert c["metodo"] == "Semelhança" and c["eaps"] == ["4.7"]


def test_excel_sem_erros(caso, tmp_path):
    d, obra, procs, mem, cz = caso
    saida = str(tmp_path / "ocs.xlsx")
    excel.gerar(obra, procs, cz, saida, mem, str(d / "compras"))
    wb = openpyxl.load_workbook(saida)
    assert wb.sheetnames == ["Resumo", "Itens 137", "Cruzamento memória", "Pastas e status"]
    ws = wb["Itens 137"]
    assert [c.value for c in ws[4]][:4] == ["Status", "Tarefa", "Nº pedido", "Nº OC"]
    rec = recalcular(saida, str(tmp_path))
    if rec:
        wr = openpyxl.load_workbook(rec, data_only=True)
        for w in wr.worksheets:
            for row in w.iter_rows(values_only=True):
                assert not any(isinstance(v, str) and v.startswith("#") for v in row), (w.title, row)
        res = {r[0]: r[1] for r in wr["Resumo"].iter_rows(min_row=6, values_only=True)}
        assert res["Cruzamento: casadas por Nº pedido"] == 1
        assert res["Pastas para verificar"] == 1
        assert res["Erros de fórmula no arquivo"] == 0
