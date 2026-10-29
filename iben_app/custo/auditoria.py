"""Auditoria fixa da Analise de Custos (parte mecanica da skill analise-custo-obra-iben).

So estrutura da planilha: valor digitado onde deveria haver formula, SUM de subtotal
desalinhado, +L (retencao) em parte das linhas, itens sem EAP no fim da Memoria, links
externos e checagens FALSO. A estimativa a 100% e julgamento e continua com o engenheiro.
O arquivo original nunca e regravado: a saida e um .xlsx novo.
"""
import datetime as dt
import os
import re

import openpyxl
from openpyxl.utils import column_index_from_string as CI
from openpyxl.utils import get_column_letter as L

from iben_app.comum import COR_ALERTA, estilo, norm, num, tabela

ABA_MEM = "Memória de Cálculo - Previsão"
COLS_MEM = ("I", "P")
LINHA0 = 8
RX_SUM = re.compile(r"^=\s*SUM\(\s*\$?([A-Z]+)\$?(\d+)\s*:\s*\$?([A-Z]+)\$?(\d+)\s*\)\s*$", re.I)


def _aba(wb, nome):
    for ws in wb.worksheets:
        if norm(ws.title) == norm(nome):
            return ws
    return None


def achado(aba, cel, tipo, texto, correcao, valor=None):
    return {"aba": aba, "celula": cel, "tipo": tipo, "achado": texto, "correcao": correcao, "valor": valor}


def memoria(wf, wv):
    """wf: workbook com formulas; wv: com valores. Retorna lista de achados."""
    ws, wsv = _aba(wf, ABA_MEM), _aba(wv, ABA_MEM)
    if ws is None:
        return [achado(ABA_MEM, "", "Estrutura", "aba não encontrada", "conferir o nome da aba")]
    out = []
    c0, c1 = CI(COLS_MEM[0]), CI(COLS_MEM[1])
    ult = ws.max_row
    eap_col = 2
    subtot = {}  # linhas de subtotal (SUM) por coluna
    for c in range(c0, c1 + 1):
        letra = L(c)
        form, const, com_l, sem_l = [], [], [], []
        for r in range(LINHA0, ult + 1):
            v = ws.cell(r, c).value
            if isinstance(v, str) and v.startswith("="):
                m = RX_SUM.match(v)
                if m:
                    subtot.setdefault(letra, []).append((r, int(m.group(2)), int(m.group(4))))
                    continue
                form.append(r)
                (com_l if re.search(rf"\+\s*\$?L\$?{r}(?!\d)", v.upper()) else sem_l).append(r)
            elif num(v):
                const.append(r)
        if form and const and len(form) >= 0.7 * (len(form) + len(const)):
            for r in const:
                out.append(achado(ws.title, f"{letra}{r}", "Valor digitado",
                                  f"número digitado numa coluna em que {len(form)} de {len(form) + len(const)} linhas são fórmula",
                                  f"refazer a fórmula da coluna {letra} na linha {r}", ws.cell(r, c).value))
        if com_l and sem_l and letra != "L":
            minoria, maioria = (com_l, "sem") if len(com_l) < len(sem_l) else (sem_l, "com")
            for r in minoria:
                out.append(achado(ws.title, f"{letra}{r}", "+L parcial",
                                  f"fórmula {'com' if maioria == 'sem' else 'sem'} +L (retenção) enquanto a maioria da coluna está {maioria}",
                                  "padronizar a fórmula da coluna (retenção já vem no incorrido)", wsv.cell(r, c).value))
    # SUM de subtotal: deve comecar na primeira linha do bloco e terminar na linha anterior
    for letra, lst in subtot.items():
        ant = LINHA0 - 1
        for r, a, b in sorted(lst):
            ini = ant + 1
            if a != ini or b != r - 1:
                out.append(achado(ws.title, f"{letra}{r}", "Subtotal desalinhado",
                                  f"SUM({letra}{a}:{letra}{b}); o bloco vai de {ini} a {r - 1}",
                                  f"=SUM({letra}{ini}:{letra}{r - 1})", wsv.cell(r, CI(letra)).value))
            ant = r
    # itens sem EAP (o fim da Memoria costuma acumular linhas soltas)
    for r in range(LINHA0, ult + 1):
        tem_valor = any(num(wsv.cell(r, c).value) and wsv.cell(r, c).value for c in range(c0 - 1, c1 + 1))
        if tem_valor and wsv.cell(r, eap_col).value in (None, "") and not any(
                isinstance(ws.cell(r, c).value, str) and RX_SUM.match(ws.cell(r, c).value) for c in range(c0, c1 + 1)):
            out.append(achado(ws.title, f"B{r}", "Sem EAP", "linha com valor e sem EAP: fica fora de qualquer EAP",
                              "classificar com o engenheiro da obra", wsv.cell(r, 8).value))
    return out


def links_externos(wf, path):
    out = []
    n = len(getattr(wf, "_external_links", []) or [])
    if n:
        out.append(achado("(arquivo)", "", "Link externo", f"{n} vínculo(s) com outras pastas de trabalho",
                          "Dados > Editar Links: confirmar a origem ou quebrar o vínculo"))
    for ws in wf.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("=") and re.search(r"\[\d+\]|\.xls[xmb]?\]", c.value, re.I):
                    out.append(achado(ws.title, c.coordinate, "Link externo", "fórmula aponta para outro arquivo",
                                      "confirmar a origem", None))
    return out


def checagens_falso(wv, wf):
    out = []
    for ws in wv.worksheets:
        wsf = wf[ws.title]
        for row in ws.iter_rows():
            for c in row:
                if c.value is False or (isinstance(c.value, str) and c.value.strip().upper() in ("FALSO", "FALSE")):
                    f = wsf[c.coordinate].value
                    if isinstance(f, str) and f.startswith("="):
                        out.append(achado(ws.title, c.coordinate, "Checagem FALSO", f"checagem da planilha deu FALSO ({f[:60]})",
                                          "achar a diferença que a checagem aponta"))
                elif isinstance(c.value, str) and c.value.startswith("#") and c.value.rstrip("!?/0A").upper() in (
                        "#REF", "#DIV", "#N", "#VALUE", "#VALOR", "#NAME", "#NOME", "#NUM", "#NULL", "#NULO"):
                    out.append(achado(ws.title, c.coordinate, "Erro de fórmula", f"célula com {c.value}", "corrigir a fórmula"))
    return out


def auditar(path):
    wf = openpyxl.load_workbook(path, data_only=False, keep_links=True)
    wv = openpyxl.load_workbook(path, data_only=True)
    ach = memoria(wf, wv) + links_externos(wf, path) + checagens_falso(wv, wf)
    return ach


def gerar_excel(path, achados, saida):
    e = estilo()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Auditoria"
    ws["A1"] = f"Auditoria | {os.path.basename(path)}"
    ws["A1"].font = e["titulo"]
    ws["A2"] = (f"Gerado em {dt.datetime.now():%d/%m/%Y %H:%M}. Só estrutura da planilha; impacto em R$ a validar "
                "com o engenheiro. O arquivo original não foi alterado.")
    ws["A2"].font = e["sub"]
    linhas = [[a["aba"], a["celula"], a["tipo"], a["achado"], a["correcao"], a["valor"] if num(a["valor"]) else None]
              for a in achados]
    tabela(ws, ["Aba", "Célula", "Tipo", "O que foi achado", "Correção sugerida", "Valor na célula (R$)"], linhas,
           [26, 9, 20, 60, 45, 18], linha0=4, formatos={5: "#,##0.00"})
    for r in range(5, 5 + len(linhas)):
        ws.cell(r, 4).alignment = ws.cell(r, 5).alignment = e["quebra"]
        if ws.cell(r, 3).value in ("Checagem FALSO", "Erro de fórmula", "Subtotal desalinhado"):
            ws.cell(r, 3).fill = e["fill"](COR_ALERTA)
    if not linhas:
        ws["A5"] = "Nenhum achado."
    wb.save(saida)
    return saida
